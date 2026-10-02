"""
Export de données vers Excel et CSV.

Ce module fournit des fonctions pour exporter les données des élèves
et des paiements vers des fichiers Excel ou CSV.
"""

import csv
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def exporter_eleves_csv(eleves, chemin_fichier: str, devise: str = "FCFA"):
    """
    Exporte la liste des élèves vers un fichier CSV.

    Args:
        eleves: Liste des élèves avec totaux
        chemin_fichier: Chemin du fichier CSV à créer
        devise: Devise pour le formatage des montants
    """
    with open(chemin_fichier, 'w', newline='', encoding='utf-8-sig') as csvfile:
        writer = csv.writer(csvfile, delimiter=';')

        # En-têtes
        writer.writerow([
            'Nom', 'Prénom', 'Classe', 'Année scolaire',
            'Total dû', 'Total payé', 'Solde', 'Statut'
        ])

        # Données
        for eleve_avec_totaux in eleves:
            eleve = eleve_avec_totaux.eleve
            solde = max(0, eleve.total_du - eleve_avec_totaux.total_paye)

            from edupaie.services.calculs import determiner_statut
            statut = determiner_statut(eleve.total_du, eleve_avec_totaux.total_paye).value

            writer.writerow([
                eleve.nom,
                eleve.prenom,
                eleve_avec_totaux.nom_classe,
                eleve.annee_scolaire,
                f"{eleve.total_du} {devise}",
                f"{eleve_avec_totaux.total_paye} {devise}",
                f"{solde} {devise}",
                statut
            ])


def exporter_eleves_excel(eleves, chemin_fichier: str, devise: str = "FCFA"):
    """
    Exporte la liste des élèves vers un fichier Excel.

    Args:
        eleves: Liste des élèves avec totaux
        chemin_fichier: Chemin du fichier Excel à créer
        devise: Devise pour le formatage des montants
    """
    if not OPENPYXL_AVAILABLE:
        raise ImportError("openpyxl n'est pas installé. Installez-le avec: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Élèves"

    # Styles
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # En-têtes
    headers = ['Nom', 'Prénom', 'Classe', 'Année scolaire', 'Total dû', 'Total payé', 'Solde', 'Statut']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Données
    for row, eleve_avec_totaux in enumerate(eleves, 2):
        eleve = eleve_avec_totaux.eleve
        solde = max(0, eleve.total_du - eleve_avec_totaux.total_paye)

        from edupaie.services.calculs import determiner_statut
        statut = determiner_statut(eleve.total_du, eleve_avec_totaux.total_paye).value

        data = [
            eleve.nom,
            eleve.prenom,
            eleve_avec_totaux.nom_classe,
            eleve.annee_scolaire,
            f"{eleve.total_du} {devise}",
            f"{eleve_avec_totaux.total_paye} {devise}",
            f"{solde} {devise}",
            statut
        ]

        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Ajuster la largeur des colonnes
    column_widths = [20, 20, 15, 15, 15, 15, 15, 15]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    wb.save(chemin_fichier)


def exporter_paiements_csv(paiements, chemin_fichier: str, devise: str = "FCFA"):
    """
    Exporte l'historique des paiements vers un fichier CSV.

    Args:
        paiements: Liste des paiements
        chemin_fichier: Chemin du fichier CSV à créer
        devise: Devise pour le formatage des montants
    """
    with open(chemin_fichier, 'w', newline='', encoding='utf-8-sig') as csvfile:
        writer = csv.writer(csvfile, delimiter=';')

        # En-têtes
        writer.writerow([
            'Numéro de reçu', 'Date', 'Heure', 'Montant', 'Mode',
            'Nom du payeur', 'État'
        ])

        # Données
        for paiement in paiements:
            from edupaie.utils.formatage import get_mode_libelle
            mode_libelle = get_mode_libelle(paiement.mode)

            writer.writerow([
                paiement.numero_recu,
                paiement.date_paiement,
                paiement.heure_paiement[:5],
                f"{paiement.montant} {devise}",
                mode_libelle,
                paiement.nom_payeur or "—",
                "Annulé" if paiement.est_annule else "Valide"
            ])


def exporter_paiements_excel(paiements, chemin_fichier: str, devise: str = "FCFA"):
    """
    Exporte l'historique des paiements vers un fichier Excel.

    Args:
        paiements: Liste des paiements
        chemin_fichier: Chemin du fichier Excel à créer
        devise: Devise pour le formatage des montants
    """
    if not OPENPYXL_AVAILABLE:
        raise ImportError("openpyxl n'est pas installé. Installez-le avec: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Paiements"

    # Styles
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # En-têtes
    headers = ['Numéro de reçu', 'Date', 'Heure', 'Montant', 'Mode', 'Nom du payeur', 'État']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Données
    for row, paiement in enumerate(paiements, 2):
        from edupaie.utils.formatage import get_mode_libelle
        mode_libelle = get_mode_libelle(paiement.mode)

        data = [
            paiement.numero_recu,
            paiement.date_paiement,
            paiement.heure_paiement[:5],
            f"{paiement.montant} {devise}",
            mode_libelle,
            paiement.nom_payeur or "—",
            "Annulé" if paiement.est_annule else "Valide"
        ]

        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Ajuster la largeur des colonnes
    column_widths = [20, 12, 10, 15, 15, 20, 10]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    wb.save(chemin_fichier)
